"""
Parsing stage. Text extraction for digital documents; OCR for scanned
PDFs and images. The image-of-a-spreadsheet edge case (Section 3.2) is
caught HERE, not at triage — via a deterministic OCR-confidence
threshold, keeping it in the pipeline rather than needing an agent
(Design Doc Section 6's boundary rule).
"""
import subprocess
import tempfile
from pathlib import Path

import pdfplumber
import pytesseract
from PIL import Image
from docx import Document as DocxDocument
import openpyxl
import csv as csv_module
from email import message_from_bytes
from bs4 import BeautifulSoup

MIN_TEXT_CHARS = 20
LOW_OCR_CONFIDENCE_QUARANTINE_THRESHOLD = 0.55  # below this, don't trust the extraction at all


class ParseResult:
    def __init__(self, text: str = "", ocr_confidence: float = None,
                 quarantine_reason: str = None, attachments: list = None):
        self.text = text
        self.ocr_confidence = ocr_confidence
        self.quarantine_reason = quarantine_reason  # set if OCR confidence too low to trust
        self.attachments = attachments or []  # list of (filename, bytes) for email attachments


def is_scanned_pdf(path: Path) -> bool:
    with pdfplumber.open(path) as pdf:
        text = "".join(p.extract_text() or "" for p in pdf.pages)
    return len(text.strip()) < MIN_TEXT_CHARS


def _ocr_image(img: Image.Image) -> tuple[str, float]:
    data = pytesseract.image_to_data(img, output_type=pytesseract.Output.DICT)
    words, confidences = [], []
    for i, word in enumerate(data["text"]):
        if word.strip():
            words.append(word)
            conf = data["conf"][i]
            if isinstance(conf, (int, float)) and conf >= 0:
                confidences.append(conf / 100.0)
    avg_conf = sum(confidences) / len(confidences) if confidences else 0.0
    return " ".join(words), avg_conf


def parse_pdf(path: Path) -> ParseResult:
    if not is_scanned_pdf(path):
        with pdfplumber.open(path) as pdf:
            text = "\n".join(p.extract_text() or "" for p in pdf.pages)
        return ParseResult(text=text)

    with tempfile.TemporaryDirectory() as tmp:
        prefix = f"{tmp}/page"
        subprocess.run(["pdftoppm", "-jpeg", "-r", "200", str(path), prefix], check=True, capture_output=True)
        all_text, confs = [], []
        for page_path in sorted(Path(tmp).glob("page-*.jpg")):
            text, conf = _ocr_image(Image.open(page_path))
            all_text.append(text)
            confs.append(conf)
        avg_conf = sum(confs) / len(confs) if confs else 0.0

    reason = "low_confidence_ocr" if avg_conf < LOW_OCR_CONFIDENCE_QUARANTINE_THRESHOLD else None
    return ParseResult(text="\n".join(all_text), ocr_confidence=avg_conf, quarantine_reason=reason)


def parse_image(path: Path) -> ParseResult:
    text, conf = _ocr_image(Image.open(path))
    reason = "low_confidence_ocr" if conf < LOW_OCR_CONFIDENCE_QUARANTINE_THRESHOLD else None
    return ParseResult(text=text, ocr_confidence=conf, quarantine_reason=reason)


def parse_docx(path: Path) -> ParseResult:
    doc = DocxDocument(str(path))
    return ParseResult(text="\n".join(p.text for p in doc.paragraphs if p.text.strip()))


def parse_xlsx(path: Path) -> ParseResult:
    wb = openpyxl.load_workbook(path, data_only=True)
    lines = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            lines.append(" | ".join(str(c) for c in row if c is not None))
    return ParseResult(text="\n".join(lines))


def parse_csv(path: Path) -> ParseResult:
    with open(path, newline="") as f:
        rows = list(csv_module.reader(f))
    return ParseResult(text="\n".join(" | ".join(r) for r in rows))


def parse_html(path: Path) -> ParseResult:
    soup = BeautifulSoup(path.read_text(encoding="utf-8"), "html.parser")
    return ParseResult(text=soup.get_text(separator="\n"))


def parse_txt(path: Path) -> ParseResult:
    return ParseResult(text=path.read_text(encoding="utf-8"))


def parse_eml(path: Path) -> ParseResult:
    msg = message_from_bytes(path.read_bytes())
    body_parts = []
    attachments = []
    for part in msg.walk():
        if part.get_filename():
            attachments.append((part.get_filename(), part.get_payload(decode=True)))
        elif part.get_content_type() == "text/plain":
            payload = part.get_payload(decode=True)
            if payload:
                body_parts.append(payload.decode(errors="replace"))
    header_text = f"From: {msg['From']}\nTo: {msg['To']}\nSubject: {msg['Subject']}\n"
    return ParseResult(text=header_text + "\n".join(body_parts), attachments=attachments)


PARSERS = {
    "pdf": parse_pdf, "docx": parse_docx, "xlsx": parse_xlsx, "csv": parse_csv,
    "html": parse_html, "htm": parse_html, "txt": parse_txt, "eml": parse_eml,
    "image": parse_image,
}


def parse_document(path: Path, fmt: str) -> ParseResult:
    parser = PARSERS.get(fmt)
    if parser is None:
        raise ValueError(f"No parser for format: {fmt}")
    return parser(path)
